"""Reads an uploaded file into pages and puts it in the index. Registered as the upload hook of the platform."""
from __future__ import annotations

import logging
from pathlib import Path

from rafeeq.platform import service

from . import store

log = logging.getLogger("rafeeq.knowledge")
TEXT = {".txt", ".md", ".csv"}


def read_pages(path: Path) -> list[str]:
    """The text of each page (a slide, or a block of a document, counts as a page). Unsupported type -> ValueError."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return [(p.extract_text() or "") for p in PdfReader(str(path)).pages]
    if ext in TEXT:
        text = path.read_text(encoding="utf-8", errors="ignore")
        words = text.split()
        return [" ".join(words[i:i + 350]) for i in range(0, len(words), 350)]
    if ext == ".docx":
        import docx
        words = " ".join(p.text for p in docx.Document(str(path)).paragraphs).split()
        return [" ".join(words[i:i + 350]) for i in range(0, len(words), 350)]
    if ext == ".pptx":
        from pptx import Presentation
        return [" ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame) for slide in Presentation(str(path)).slides]
    raise ValueError(f"cannot read {ext or 'this'} files yet")


def display_name(path: Path) -> str:
    """The platform saves files as '<file_id>__<name>': the teacher knows the file by <name>."""
    return path.name.split("__", 1)[1] if "__" in path.name else path.name


def index_file(room_id: str, file_id: str, path: Path | str) -> int:
    path = Path(path)
    s = service.get_store()
    s.set_file_status(room_id, file_id, "indexing")
    try:
        n = store.add_file(room_id, file_id, display_name(path), read_pages(path))
    except Exception as e:                                     # noqa: BLE001 - one bad file must not break the upload
        log.warning("could not index %s: %s", path.name, e)
        s.set_file_status(room_id, file_id, "failed")
        return 0
    s.set_file_status(room_id, file_id, "ready" if n else "failed")      # a file with no readable text is not "ready"
    return n


def reindex_pending() -> int:
    """At start-up: files uploaded before the index existed (or while the app was down) get indexed now."""
    s, done = service.get_store(), 0
    for room_id, files in list(s.db.get("files", {}).items()):
        for f in files:
            path = Path(f.get("path", ""))
            if f.get("status") in ("uploaded", "indexing") and path.is_file() and not store.indexed(room_id, f["file_id"]):
                done += bool(index_file(room_id, f["file_id"], path))
    return done


def install() -> None:
    service.on_upload(index_file)
