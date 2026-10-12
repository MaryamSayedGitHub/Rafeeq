"""The teacher studios as separate, connected pages: Presentation Studio (decks, themes, AI pictures),
Visual & Media Studio, My Learning Platform and the library that ties them together.

Runs against the REAL app (real routes, real stores, real files in a temp folder). Providers that cost
money or need the internet (Wan picture model, edge-tts, YouTube thumbnails) are replaced by small fakes.
    uv run pytest tests/test_studios.py
"""
import io
import re
import zipfile
from types import SimpleNamespace

import pytest

MONA = {"Authorization": "Bearer demo-token-mona-456"}   # teacher u_mona
SARA = {"Authorization": "Bearer demo-token-sara-123"}   # student u_sara
HX = {"HX-Request": "true"}


def png_bytes(w=240, h=120) -> bytes:
    cv2 = pytest.importorskip("cv2")
    np = pytest.importorskip("numpy")
    img = np.full((h, w, 3), 200, np.uint8)
    cv2.rectangle(img, (10, 10), (60, 60), (0, 0, 255), -1)
    return cv2.imencode(".png", img)[1].tobytes()


class FakeStore:
    """Same calls as the LangGraph store the app uses (put / get / search / delete)."""

    def __init__(self):
        self.data = {}

    def put(self, namespace, key, value):
        self.data[(tuple(namespace), key)] = dict(value)

    def get(self, namespace, key):
        value = self.data.get((tuple(namespace), key))
        return SimpleNamespace(key=key, value=value) if value is not None else None

    def search(self, namespace, limit=10):
        return [SimpleNamespace(key=k, value=v) for (ns, k), v in self.data.items() if ns == tuple(namespace)][:limit]

    def delete(self, namespace, key):
        self.data.pop((tuple(namespace), key), None)


@pytest.fixture
def folders(tmp_path, monkeypatch):
    for name in ("ASSET_DIR", "OUTPUT_DIR"):
        folder = (tmp_path / name).resolve()
        folder.mkdir()
        monkeypatch.setattr(f"rafeeq.teacher.lecturer.settings.{name}", folder)
    return tmp_path


@pytest.fixture
def store(folders):
    return FakeStore()


@pytest.fixture
def client(monkeypatch, folders):
    pytest.importorskip("langgraph")
    from fastapi.testclient import TestClient

    from rafeeq.platform import service as platform

    monkeypatch.setattr("rafeeq.teacher.lecturer.settings.DATABASE_URL", None)   # in-memory store
    monkeypatch.setenv("RAFEEQ_WEB_SECRET", "test-secret")
    avatars = (folders / "AVATARS").resolve()
    avatars.mkdir()
    monkeypatch.setattr("rafeeq.teacher.routes.AVATAR_DIR", avatars)
    platform.set_store(platform.PlatformStore(folders / "platform", seed_demo=True))
    from rafeeq.main import app

    with TestClient(app) as c:
        yield c
    platform.set_store(None)


def teacher_store(client):
    return client.app.state.teacher.store


def new_account(client, role, email):
    res = client.post("/auth/register", json={"name": f"{role} {email}", "email": email, "password": "secret1", "role": role})
    assert res.status_code in (200, 201), res.text
    return {"Authorization": f"Bearer {res.json()['token']}"}


def upload(client, headers=MONA, name="p.png", data=None, ctype="image/png", **form):
    res = client.post("/teacher/media/upload", headers=headers, files={"file": (name, data or png_bytes(), ctype)}, data=form)
    assert res.status_code == 201, res.text
    return res.json()


def make_deck(client, owner="u_mona", session_id="s_1", room_id=""):
    from rafeeq.teacher import decks

    return decks.create_deck(
        teacher_store(client), owner, title="الخلية", course_name="Biology", lecturer="Mona", room_id=room_id, session_id=session_id,
        slides=[{"title": "مقدمة", "bullets": ["أ", "ب"], "notes": "ن"}, {"title": "الخاتمة", "bullets": ["ج"]}])


# =============================================================== pure modules
def test_themes_validate_colours_and_extract_from_a_real_pptx():
    from rafeeq.teacher import themes

    ok = themes.normalize({"name": " Mine ", "font": "Cairo", "colors": {"blue": "#abcdef"}})
    assert ok["colors"]["blue"] == "#ABCDEF" and ok["colors"]["ink"] == themes.PRESETS["rafeeq"]["colors"]["ink"]
    for bad in ({"colors": {"blue": "red"}}, {"colors": {"ink": "#12345"}}, "nope"):
        with pytest.raises(themes.ThemeError):
            themes.normalize(bad)

    pptx = pytest.importorskip("pptx")
    buf = io.BytesIO()
    pptx.Presentation().save(buf)
    got = themes.extract_from_pptx(buf.getvalue(), fallback_name="Default")
    assert set(got["colors"]) == set(themes.COLOR_KEYS) and all(re.fullmatch(r"#[0-9A-F]{6}", c) for c in got["colors"].values())
    with pytest.raises(themes.ThemeError):
        themes.extract_from_pptx(b"not a zip")
    empty = io.BytesIO()
    with zipfile.ZipFile(empty, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(themes.ThemeError):
        themes.extract_from_pptx(empty.getvalue())


def test_image_edits_change_the_picture_and_refuse_bad_values():
    cv2 = pytest.importorskip("cv2")
    np = pytest.importorskip("numpy")
    from rafeeq.teacher import media_tools

    def shape(data):
        return cv2.imdecode(np.frombuffer(data, dtype=np.uint8), 1).shape[:2]

    src = png_bytes(240, 120)
    assert shape(media_tools.edit_image(src, ".png", {"rotate": 90})) == (240, 120)
    assert shape(media_tools.edit_image(src, ".png", {"width": 120})) == (60, 120)
    assert shape(media_tools.edit_image(src, ".png", {"crop": {"x": 0, "y": 0, "w": 50, "h": 50}})) == (60, 120)
    gray = cv2.imdecode(np.frombuffer(media_tools.edit_image(src, ".png", {"grayscale": True}), dtype=np.uint8), 1)
    assert (gray[..., 0] == gray[..., 1]).all()
    for bad in ({"rotate": 45}, {"flip": "x"}, {"crop": {"x": 90, "y": 0, "w": 50, "h": 50}}, {"brightness": 500}):
        with pytest.raises(media_tools.MediaError):
            media_tools.edit_image(src, ".png", bad)
    with pytest.raises(media_tools.MediaError):
        media_tools.edit_image(b"not an image", ".png", {})
    assert [media_tools.kind_for(n) for n in ("a.JPG", "b.mp4", "c.mp3", "d.exe")] == ["image", "video", "audio", ""]


def test_decks_are_per_teacher_idempotent_and_only_accept_own_images(store):
    from rafeeq.teacher import assets, decks

    slides = [{"title": "One", "bullets": ["a"]}]
    d = decks.create_deck(store, "u_mona", title="T", slides=slides, session_id="s1")
    assert decks.create_deck(store, "u_mona", title="T", slides=slides, session_id="s1")["deck_id"] == d["deck_id"]
    assert decks.get_deck(store, "u_other", d["deck_id"]) is None and decks.list_decks(store, "u_other") == []
    with pytest.raises(decks.DeckError):
        decks.create_deck(store, "u_mona", title="T", slides=[{"title": "", "bullets": []}])
    with pytest.raises(decks.DeckError):
        decks.update_deck(store, "u_mona", d["deck_id"], theme_id="no-such-theme")
    other_image = assets.register_asset(store, "u_other", kind="image", title="x", origin="upload")
    with pytest.raises(decks.DeckError):
        decks.set_slide_image(store, "u_mona", d["deck_id"], 0, other_image["asset_id"])
    with pytest.raises(decks.DeckError):
        decks.set_slide_image(store, "u_mona", d["deck_id"], 5, "")
    mine = assets.register_asset(store, "u_mona", kind="image", title="mine", origin="upload")
    assert decks.set_slide_image(store, "u_mona", d["deck_id"], 0, mine["asset_id"])["slides"][0]["image_asset_id"] == mine["asset_id"]
    assert decks.delete_deck(store, "u_mona", d["deck_id"]) and decks.get_deck(store, "u_mona", d["deck_id"]) is None


def test_courses_snapshot_only_what_is_ready_and_checks_ownership(store):
    from rafeeq.teacher import assets, courses

    ready = assets.register_asset(store, "u_mona", kind="source_video", title="ok", origin="external", source_url="https://youtu.be/x")
    pending = assets.register_asset(store, "u_mona", kind="image", title="later", origin="upload", status="pending")
    foreign = assets.register_asset(store, "u_other", kind="image", title="theirs", origin="upload")
    c = courses.create_course(store, "u_mona", "Biology")
    with pytest.raises(courses.CourseError):
        courses.create_course(store, "u_mona", "x")
    with pytest.raises(courses.CourseError):
        courses.add_lesson(store, "u_mona", c["course_id"], "L", asset_ids=[foreign["asset_id"]])
    with pytest.raises(courses.CourseError):
        courses.publish(store, "u_mona", c["course_id"], "r1")            # no lessons yet
    c = courses.add_lesson(store, "u_mona", c["course_id"], "Lesson 1", asset_ids=[ready["asset_id"], pending["asset_id"]])
    lesson = c["lessons"][0]["lesson_id"]
    published = courses.publish(store, "u_mona", c["course_id"], "r1")
    assert published["published_room_id"] == "r1" and not courses.has_unpublished_changes(published)
    snap = courses.published_for_room(store, "r1")
    assert [a["asset_id"] for a in snap[0]["lessons"][0]["assets"]] == [ready["asset_id"]]   # the pending one is left out
    assert courses.published_owner(store, "r1", c["course_id"], ready["asset_id"]) == "u_mona"
    assert courses.published_owner(store, "r1", c["course_id"], pending["asset_id"]) == ""
    assert courses.published_owner(store, "r2", c["course_id"], ready["asset_id"]) == ""
    edited = courses.update_lesson(store, "u_mona", c["course_id"], lesson, title="Renamed")
    assert courses.has_unpublished_changes(edited)
    courses.unpublish(store, "u_mona", c["course_id"])
    assert courses.published_for_room(store, "r1") == []
    assert courses.delete_course(store, "u_mona", c["course_id"]) and courses.get_course(store, "u_mona", c["course_id"]) is None


# =============================================================== API: themes, decks, export
def test_the_studio_api_needs_a_teacher_token(client):
    for path in ("/teacher/decks", "/teacher/themes", "/teacher/courses", "/teacher/media/capabilities"):
        assert client.get(path).status_code == 401, path
        assert client.get(path, headers=SARA).status_code == 403, path


def test_extract_a_theme_from_a_pptx_and_list_it(client):
    pptx = pytest.importorskip("pptx")
    buf = io.BytesIO()
    pptx.Presentation().save(buf)
    files = {"file": ("brand.pptx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    res = client.post("/teacher/themes/extract", headers=MONA, files=files, data={"name": "My brand"})
    assert res.status_code == 201 and res.json()["kind"] == "theme" and res.json()["theme"]["name"] == "My brand"
    again = client.post("/teacher/themes/extract", headers=MONA, files=files, data={"name": "My brand"})
    assert again.json()["asset_id"] == res.json()["asset_id"]                              # same file, no duplicate
    listing = client.get("/teacher/themes", headers=MONA).json()
    assert [p["theme_id"] for p in listing["presets"]][0].startswith("preset:")
    assert [c["theme_id"] for c in listing["custom"]] == [res.json()["asset_id"]]
    assert client.get("/teacher/themes", headers=new_account(client, "teacher", "t2@x.com")).json()["custom"] == []
    bad = client.post("/teacher/themes/extract", headers=MONA, files={"file": ("x.pptx", b"nope", "application/zip")})
    assert bad.status_code == 422


def test_edit_a_deck_and_export_it_with_a_theme_and_a_picture(client):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = make_deck(client)
    other = new_account(client, "teacher", "t2@x.com")
    assert client.get(f"/teacher/decks/{deck['deck_id']}", headers=other).status_code == 404       # not theirs
    assert client.get("/teacher/decks", headers=other).json()["decks"] == []

    listing = client.get("/teacher/decks", headers=MONA).json()["decks"]
    assert [d["deck_id"] for d in listing] == [deck["deck_id"]] and listing[0]["slide_count"] == 2

    from rafeeq.teacher import assets, themes
    custom = themes.normalize({"name": "Green", "colors": {"paper": "#EAF4EE", "blue": "#087F73"}})
    theme_asset = assets.register_asset(teacher_store(client), "u_mona", kind="theme", title="Green", origin="upload", settings_used=custom)
    picture = upload(client)

    slides = [{"title": "مقدمة معدّلة", "bullets": ["نقطة 1", "نقطة 2"], "notes": "ملاحظة", "image_asset_id": picture["asset_id"]},
              {"title": "الخاتمة", "bullets": ["ج"]}]
    saved = client.put(f"/teacher/decks/{deck['deck_id']}", headers=MONA,
                       json={"title": "الخلية 2", "slides": slides, "theme_id": theme_asset["asset_id"]})
    assert saved.status_code == 200, saved.text
    view = saved.json()
    assert view["title"] == "الخلية 2" and view["theme_name"] == "Green" and view["theme"]["colors"]["paper"] == "#EAF4EE"
    assert view["slides"][0]["image_file_url"] == f"/teacher/assets/{picture['asset_id']}/file" and view["slides"][1]["image_file_url"] is None

    # refusals
    assert client.put(f"/teacher/decks/{deck['deck_id']}", headers=MONA, json={"theme_id": "preset:nope"}).status_code == 422
    assert client.put(f"/teacher/decks/{deck['deck_id']}", headers=MONA, json={"slides": []}).status_code == 422
    foreign = upload(client, headers=other)
    assert client.post(f"/teacher/decks/{deck['deck_id']}/slides/1/image", headers=MONA,
                       json={"asset_id": foreign["asset_id"]}).status_code == 422                      # someone else's image
    assert client.post(f"/teacher/decks/{deck['deck_id']}/slides/9/image", headers=MONA, json={"asset_id": ""}).status_code == 422

    exported = client.post(f"/teacher/decks/{deck['deck_id']}/export", headers=MONA)
    assert exported.status_code == 200, exported.text
    asset = exported.json()
    assert asset["kind"] == "pptx" and asset["origin"] == "edited" and asset["status"] == "ready" and asset["file_url"]
    data = client.get(asset["file_url"], headers=MONA).content
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) >= 2
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sl in prs.slides for sh in sl.shapes)       # the picture is in the file
    fills = set()
    for sl in prs.slides:
        for sh in sl.shapes:
            try:
                if sh.fill.type == 1:
                    fills.add(str(sh.fill.fore_color.rgb))
            except (AttributeError, TypeError, ValueError):
                pass
    assert "EAF4EE" in fills                                                                          # the theme's paper colour

    again = client.post(f"/teacher/decks/{deck['deck_id']}/export", headers=MONA).json()
    assert again["version"] == asset["version"] + 1 and again["parent_asset_id"] == asset["asset_id"]  # versions, not overwrites
    assert client.get(asset["file_url"], headers=MONA).content == data                                 # the first file is untouched
    assert client.get(f"/teacher/decks/{deck['deck_id']}", headers=MONA).json()["exports"] == 2


# =============================================================== API: AI pictures for slides
def fake_picture_provider(monkeypatch, status="SUCCEEDED"):
    calls = SimpleNamespace(submitted=[], polled=[])
    monkeypatch.setattr("rafeeq.teacher.visuals.suggest_prompt", lambda title, bullets: (f"diagram of {title}", True))

    def submit(prompt):
        calls.submitted.append(prompt)
        return {"ok": True, "task_id": f"T{len(calls.submitted)}", "model": "wan-test"}

    def poll(task_id):
        calls.polled.append(task_id)
        return {"status": status, "image_url": "https://example.com/x.png" if status == "SUCCEEDED" else None,
                "error": None if status == "SUCCEEDED" else "blocked"}

    monkeypatch.setattr("rafeeq.teacher.lecturer.image_gen.submit_task", submit)
    monkeypatch.setattr("rafeeq.teacher.lecturer.image_gen.poll_task", poll)
    monkeypatch.setattr("rafeeq.teacher.lecturer.image_gen.download_image", lambda url: png_bytes())
    return calls


def test_a_generated_slide_picture_lands_in_the_library_and_on_the_slide(client, monkeypatch):
    calls = fake_picture_provider(monkeypatch)
    deck = make_deck(client)
    url = f"/teacher/decks/{deck['deck_id']}/slides/0/visual"
    started = client.post(url, headers=MONA, json={})
    assert started.status_code == 200, started.text
    job = started.json()
    assert job["task_id"] == "T1" and job["prompt"] == "diagram of مقدمة" and job["used_llm"] is True
    assert client.post(url, headers=MONA, json={}).json()["asset_id"] == job["asset_id"] and len(calls.submitted) == 1   # a retry costs nothing

    status = client.get("/teacher/visuals/T1", headers=MONA).json()
    assert status["status"] == "SUCCEEDED" and status["file_url"] and status["asset_id"] == job["asset_id"]
    assert client.get("/teacher/visuals/T1", headers=MONA).json()["status"] == "SUCCEEDED" and calls.polled == ["T1"]    # not polled again
    asset = client.get(f"/teacher/assets/{job['asset_id']}", headers=MONA).json()
    assert asset["kind"] == "image" and asset["origin"] == "ai_generated" and asset["status"] == "ready" and asset["slide_index"] == 0
    assert client.get(f"/teacher/decks/{deck['deck_id']}", headers=MONA).json()["slides"][0]["image_asset_id"] == job["asset_id"]
    assert client.get("/teacher/visuals/T1", headers=new_account(client, "teacher", "t2@x.com")).status_code == 404


def test_a_failed_picture_job_is_reported_and_recorded(client, monkeypatch):
    fake_picture_provider(monkeypatch, status="FAILED")
    deck = make_deck(client)
    job = client.post(f"/teacher/decks/{deck['deck_id']}/slides/1/visual", headers=MONA, json={"prompt": "a red cell"}).json()
    assert job["prompt"] == "a red cell" and job["used_llm"] is False
    status = client.get(f"/teacher/visuals/{job['task_id']}", headers=MONA).json()
    assert status["status"] == "FAILED" and status["error"] == "blocked"
    assert client.get(f"/teacher/assets/{job['asset_id']}", headers=MONA).json()["status"] == "failed"
    assert client.post(f"/teacher/decks/{deck['deck_id']}/slides/7/visual", headers=MONA, json={}).status_code == 404


def test_picture_generation_without_a_provider_key_is_a_clear_503(client, monkeypatch):
    monkeypatch.setattr("rafeeq.teacher.visuals.suggest_prompt", lambda t, b: ("p", False))
    monkeypatch.setattr("rafeeq.teacher.lecturer.image_gen.submit_task",
                        lambda prompt: {"ok": False, "error": "DASHSCOPE_API_KEY is not set in .env"})
    deck = make_deck(client)
    res = client.post(f"/teacher/decks/{deck['deck_id']}/slides/0/visual", headers=MONA, json={})
    assert res.status_code == 503 and "DASHSCOPE_API_KEY" in res.json()["detail"]


# =============================================================== API: media
def test_upload_and_edit_keep_the_original_and_create_versions(client):
    cv2 = pytest.importorskip("cv2")
    np = pytest.importorskip("numpy")
    original = png_bytes()
    first = upload(client, data=original, title="Cell photo")
    assert first["kind"] == "image" and first["origin"] == "upload" and first["status"] == "ready" and first["title"] == "Cell photo"
    assert upload(client, data=original)["asset_id"] == first["asset_id"]                       # same file twice = one record

    edited = client.post(f"/teacher/media/{first['asset_id']}/edit", headers=MONA, json={"ops": {"rotate": 90, "grayscale": True}})
    assert edited.status_code == 201, edited.text
    new = edited.json()
    assert new["asset_id"] != first["asset_id"] and new["version"] == 2 and new["parent_asset_id"] == first["asset_id"]
    assert new["origin"] == "edited" and new["status"] == "ready"
    pic = cv2.imdecode(np.frombuffer(client.get(new["file_url"], headers=MONA).content, dtype=np.uint8), 1)
    assert pic.shape[:2] == (240, 120)
    assert client.get(first["file_url"], headers=MONA).content == original                       # original untouched
    again = client.post(f"/teacher/media/{first['asset_id']}/edit", headers=MONA, json={"ops": {"rotate": 90, "grayscale": True}})
    assert again.json()["asset_id"] == new["asset_id"]                                           # a double click makes no duplicate

    for ops in ({"rotate": 45}, {"crop": {"x": 90, "y": 0, "w": 50, "h": 50}}):
        assert client.post(f"/teacher/media/{first['asset_id']}/edit", headers=MONA, json={"ops": ops}).status_code == 422
    other = new_account(client, "teacher", "t2@x.com")
    assert client.post(f"/teacher/media/{first['asset_id']}/edit", headers=other, json={"ops": {"rotate": 90}}).status_code == 404


def test_upload_refuses_unknown_types_and_empty_files(client):
    res = client.post("/teacher/media/upload", headers=MONA, files={"file": ("run.exe", b"MZ", "application/octet-stream")})
    assert res.status_code == 422
    res = client.post("/teacher/media/upload", headers=MONA, files={"file": ("empty.png", b"", "image/png")})
    assert res.status_code in (413, 422)


def test_video_editing_says_so_when_ffmpeg_is_missing(client, monkeypatch):
    monkeypatch.setattr("rafeeq.teacher.media_tools.ffmpeg_path", lambda: None)
    clip = upload(client, name="clip.mp4", data=b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 64, ctype="video/mp4")
    assert clip["kind"] == "video"
    res = client.post(f"/teacher/media/{clip['asset_id']}/edit", headers=MONA, json={"ops": {"trim": {"start": 0, "end": 2}}})
    assert res.status_code == 503 and "ffmpeg" in res.json()["detail"]
    assert client.get("/teacher/media/capabilities", headers=MONA).json()["ffmpeg"] is False


def test_narration_becomes_an_audio_asset(client, monkeypatch):
    def fake_speech(text, voice, dst):
        dst.write_bytes(b"ID3-fake-mp3")
        return dst

    from rafeeq.teacher import media_tools

    real_speech = media_tools.synthesize_speech
    monkeypatch.setattr(media_tools, "synthesize_speech", fake_speech)
    body = {"text": "مرحبًا بكم في درس الخلية", "voice": "ar-EG-SalmaNeural", "title": "مقدمة صوتية"}
    res = client.post("/teacher/media/speech", headers=MONA, json=body)
    assert res.status_code == 201, res.text
    audio = res.json()
    assert audio["kind"] == "audio" and audio["origin"] == "ai_generated" and audio["status"] == "ready"
    assert client.get(audio["file_url"], headers=MONA).content == b"ID3-fake-mp3"
    assert client.post("/teacher/media/speech", headers=MONA, json=body).json()["asset_id"] == audio["asset_id"]
    monkeypatch.setattr(media_tools, "synthesize_speech", real_speech)
    assert client.post("/teacher/media/speech", headers=MONA, json={**body, "voice": "xx"}).status_code == 422


def test_youtube_cover_is_saved_from_youtube_only(client, monkeypatch):
    seen = []

    class Fake:
        def __init__(self, data):
            self.data = data

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self, n=-1):
            return self.data if n < 0 else self.data[:n]

    def fake_open(req, timeout=0):
        seen.append(req.full_url)
        return Fake(png_bytes())

    monkeypatch.setattr("urllib.request.urlopen", fake_open)
    res = client.post("/teacher/media/thumbnail", headers=MONA, json={"video_id": "dQw4w9WgXcQ", "title": "Cells", "start_sec": 75})
    assert res.status_code == 201, res.text
    cover = res.json()
    assert cover["kind"] == "image" and cover["origin"] == "external" and cover["status"] == "ready"
    assert cover["source_url"] == "https://www.youtube.com/watch?v=dQw4w9WgXcQ&t=75s"
    assert seen == ["https://i.ytimg.com/vi/dQw4w9WgXcQ/hqdefault.jpg"]
    assert client.post("/teacher/media/thumbnail", headers=MONA, json={"video_id": "dQw4w9WgXcQ"}).json()["asset_id"] == cover["asset_id"]
    assert len(seen) == 1
    for bad in ("../../etc/passwd", "short", "https://evil.example/x.jpg"):
        assert client.post("/teacher/media/thumbnail", headers=MONA, json={"video_id": bad}).status_code == 422
    assert len(seen) == 1


# =============================================================== API: courses and what students see
def room_with_student(client):
    made = client.post("/rooms", headers=MONA, json={"name": "Biology 10"}).json()
    room_id, code = made["room"]["room_id"], made["join_code"]
    assert client.post("/rooms/join", headers=SARA, json={"code": code}).status_code == 200
    return room_id


def test_courses_are_published_to_the_room_and_students_see_only_the_snapshot(client):
    room_id = room_with_student(client)
    image, later = upload(client), upload(client, data=png_bytes(300, 150))
    course = client.post("/teacher/courses", headers=MONA, json={"title": "الأحياء", "room_id": room_id}).json()
    cid = course["course_id"]
    assert client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": room_id}).status_code == 422   # empty course
    with_lesson = client.post(f"/teacher/courses/{cid}/lessons", headers=MONA,
                              json={"title": "الدرس الأول", "notes": "ابدأ هنا", "asset_ids": [image["asset_id"]]})
    assert with_lesson.status_code == 201 and with_lesson.json()["lessons"][0]["assets"][0]["asset_id"] == image["asset_id"]
    lid = with_lesson.json()["lessons"][0]["lesson_id"]

    assert client.get(f"/rooms/{room_id}/courses", headers=SARA).json() == {"courses": []}                                # not published yet
    published = client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": room_id})
    assert published.status_code == 200 and published.json()["published_room_id"] == room_id
    assert published.json()["unpublished_changes"] is False

    seen = client.get(f"/rooms/{room_id}/courses", headers=SARA).json()["courses"]
    assert [c["title"] for c in seen] == ["الأحياء"] and seen[0]["lessons"][0]["notes"] == "ابدأ هنا"
    asset = seen[0]["lessons"][0]["assets"][0]
    assert asset["file_url"] == f"/rooms/{room_id}/courses/{cid}/assets/{image['asset_id']}/file"
    got = client.get(asset["file_url"], headers=SARA)
    assert got.status_code == 200 and got.content == client.get(image["file_url"], headers=MONA).content

    # edits after publishing are invisible to students until the teacher publishes again
    after = client.post(f"/teacher/courses/{cid}/lessons/{lid}/assets", headers=MONA, json={"add": later["asset_id"]}).json()
    assert after["unpublished_changes"] is True
    assert client.get(f"/rooms/{room_id}/courses/{cid}/assets/{later['asset_id']}/file", headers=SARA).status_code == 404
    client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": room_id})
    assert client.get(f"/rooms/{room_id}/courses/{cid}/assets/{later['asset_id']}/file", headers=SARA).status_code == 200

    client.post(f"/teacher/courses/{cid}/unpublish", headers=MONA)
    assert client.get(f"/rooms/{room_id}/courses", headers=SARA).json() == {"courses": []}
    assert client.get(asset["file_url"], headers=SARA).status_code == 404


def test_course_access_rules(client):
    room_id = room_with_student(client)
    other_teacher = new_account(client, "teacher", "t2@x.com")
    outsider = new_account(client, "student", "s2@x.com")
    mine, theirs = upload(client), upload(client, headers=other_teacher, data=png_bytes(50, 50))
    cid = client.post("/teacher/courses", headers=MONA, json={"title": "Course"}).json()["course_id"]
    client.post(f"/teacher/courses/{cid}/lessons", headers=MONA, json={"title": "L1", "asset_ids": [mine["asset_id"]]})

    assert client.post(f"/teacher/courses/{cid}/lessons", headers=MONA,
                       json={"title": "L2", "asset_ids": [theirs["asset_id"]]}).status_code == 422       # someone else's file
    assert client.get(f"/teacher/courses/{cid}", headers=other_teacher).status_code == 404
    assert client.get("/teacher/courses", headers=other_teacher).json()["courses"] == []
    foreign_room = client.post("/rooms", headers=other_teacher, json={"name": "Theirs"}).json()["room"]["room_id"]
    assert client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": foreign_room}).status_code == 403
    client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": room_id})
    assert client.get(f"/rooms/{room_id}/courses", headers=outsider).status_code == 403                    # not in the room
    assert client.get(f"/rooms/{room_id}/courses").status_code == 401
    assert client.get("/teacher/courses", headers=SARA).status_code == 403                                  # students use the room routes


def test_lessons_can_be_reordered_and_removed(client):
    cid = client.post("/teacher/courses", headers=MONA, json={"title": "Order"}).json()["course_id"]
    for name in ("A", "B", "C"):
        client.post(f"/teacher/courses/{cid}/lessons", headers=MONA, json={"title": f"Lesson {name}"})
    lessons = client.get(f"/teacher/courses/{cid}", headers=MONA).json()["lessons"]
    first = lessons[0]["lesson_id"]
    moved = client.post(f"/teacher/courses/{cid}/lessons/{first}/move", headers=MONA, json={"direction": "down"}).json()
    assert [l["title"] for l in moved["lessons"]] == ["Lesson B", "Lesson A", "Lesson C"]
    assert client.post(f"/teacher/courses/{cid}/lessons/{first}/move", headers=MONA, json={"direction": "sideways"}).status_code == 422
    left = client.delete(f"/teacher/courses/{cid}/lessons/{first}", headers=MONA).json()
    assert [l["title"] for l in left["lessons"]] == ["Lesson B", "Lesson C"]
    assert client.delete(f"/teacher/courses/{cid}", headers=MONA).status_code == 204
    assert client.get(f"/teacher/courses/{cid}", headers=MONA).status_code == 404


# =============================================================== the web pages, end to end
def login(app, email, password="rafeeq-demo"):
    from starlette.testclient import TestClient

    t = TestClient(app)
    page = t.get("/app/login")
    r = t.post("/app/login", data={"csrf_token": csrf_of(page.text), "email": email, "password": password})
    assert r.status_code == 200 and "/login" not in str(r.url), r.text[:300]
    return t


def csrf_of(html):
    m = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert m, "no CSRF token on the page"
    return m.group(1)


def test_each_studio_is_its_own_page_and_the_sidebar_links_them(client):
    t = login(client.app, "mona@rafeeq.demo")
    pages = {"/app/video": "اكتشاف الفيديو", "/app/presentations": "العروض التقديمية", "/app/media": "الصور والوسائط",
             "/app/learning": "منصتي التعليمية", "/app/library": "مكتبتي"}
    for path, heading in pages.items():
        r = t.get(path)
        assert r.status_code == 200 and heading in r.text, path
        assert r.text.count('aria-current="page"') >= 1, path
        for other in pages:                                   # every page carries the whole navigation
            assert f'href="{other}' in r.text, (path, other)
    assert t.get("/app/studio").status_code == 200           # no rooms yet: the rooms page with a hint, never an error
    student = login(client.app, "sara@rafeeq.demo")
    for path in pages:
        assert student.get(path).status_code == 403, path    # students never see the teacher studios
    assert 'href="/app/video' not in student.get("/app/rooms").text


def test_work_flows_between_the_studios_and_persists(client, monkeypatch):
    t = login(client.app, "mona@rafeeq.demo")
    page = t.get("/app/rooms")
    made = t.post("/app/rooms/create", data={"csrf_token": csrf_of(page.text), "name": "Biology 10"})
    room_id = re.search(r"/app/rooms/(r_[0-9a-f]+)/", str(made.url)).group(1)
    code = re.search(r'join-code-value" dir="ltr">([A-Z0-9-]+)<', made.text).group(1)
    deck = make_deck(client, room_id=room_id)
    deck_id = deck["deck_id"]
    studio = t.get(f"/app/rooms/{room_id}/studio")
    assert "ابحث عن شرح بالفيديو" not in studio.text and f"/app/video?room={room_id}" in studio.text   # Lesson Studio stays focused

    # Presentation Studio: list -> editor -> save text
    assert deck_id in t.get(f"/app/presentations?room={room_id}").text
    editor = t.get(f"/app/presentations/{deck_id}")
    assert editor.status_code == 200 and 'name="title_0"' in editor.text and "مقدمة" in editor.text
    token = csrf_of(editor.text)
    saved = t.post(f"/app/presentations/{deck_id}", data={
        "csrf_token": token, "title": "الخلية (منقّح)", "theme_id": "preset:teal", "title_0": "مقدمة جديدة",
        "bullets_0": "نقطة أولى\nنقطة ثانية\n", "notes_0": "ملاحظة", "title_1": "الخاتمة", "bullets_1": "ج", "notes_1": ""})
    assert "حُفظ العرض" in saved.text
    stored = client.get(f"/teacher/decks/{deck_id}", headers=MONA).json()
    assert stored["title"] == "الخلية (منقّح)" and stored["theme_id"] == "preset:teal"
    assert stored["slides"][0]["title"] == "مقدمة جديدة" and stored["slides"][0]["bullets"] == ["نقطة أولى", "نقطة ثانية"]

    # Media Studio: upload -> edit creates a version -> the picture goes on a slide of the deck
    media_token = csrf_of(t.get("/app/media").text)
    up = t.post("/app/media/upload", data={"csrf_token": media_token, "title": "صورة الخلية"},
                files={"file": ("cell.png", png_bytes(), "image/png")})
    asset_id = re.search(r"asset=(a_[0-9a-f]+)", str(up.url)).group(1)
    assert "تم الرفع" in up.text and f"/app/media/{asset_id}/edit" in up.text
    edit = t.post(f"/app/media/{asset_id}/edit", data={"csrf_token": media_token, "rotate": "90", "title": "صورة مدوّرة"})
    version_id = re.search(r"asset=(a_[0-9a-f]+)", str(edit.url)).group(1)
    assert version_id != asset_id and "حُفظت النسخة الجديدة" in edit.text
    nothing = t.post(f"/app/media/{asset_id}/edit", data={"csrf_token": media_token})
    assert "لم تغيّر شيئًا" in nothing.text

    picker = t.get(f"/app/presentations/{deck_id}?asset={version_id}")
    assert "صورة مدوّرة" in picker.text                                                              # the library image is offered
    attached = t.post(f"/app/presentations/{deck_id}/slides/0/image", data={"csrf_token": media_token, "asset_id": version_id})
    assert "تم تحديث صورة الشريحة" in attached.text and f"/app/assets/{version_id}/file" in attached.text
    exported = t.post(f"/app/presentations/{deck_id}/export", data={"csrf_token": media_token})
    assert "اكتمل التصدير" in exported.text and "/file?download=1" in exported.text

    # My Learning Platform: course -> lesson with library files -> publish -> the student sees it
    created = t.post("/app/learning/courses", data={"csrf_token": media_token, "title": "دورة الأحياء", "room_id": room_id})
    cid = re.search(r"/learning/(c_[0-9a-f]+)", str(created.url)).group(1)
    pending_add = t.get(f"/app/learning/{cid}?asset={version_id}")
    assert "صورة مدوّرة" in pending_add.text
    t.post(f"/app/learning/{cid}/lessons", data={"csrf_token": media_token, "title": "الدرس الأول", "notes": "اقرأ أولًا"})
    lesson_id = re.search(r'id="lesson-(l_[0-9a-f]+)"', t.get(f"/app/learning/{cid}").text).group(1)
    t.post(f"/app/learning/{cid}/lessons/{lesson_id}/assets", data={"csrf_token": media_token, "add": version_id})
    published = t.post(f"/app/learning/{cid}/publish", data={"csrf_token": media_token, "room_id": room_id})
    assert "نُشرت الدورة" in published.text

    s = login(client.app, "sara@rafeeq.demo")
    s.post("/app/rooms/join", data={"csrf_token": csrf_of(s.get("/app/rooms").text), "code": code})
    learn = s.get(f"/app/rooms/{room_id}/learn")
    assert learn.status_code == 200 and "دورة الأحياء" in learn.text and "الدرس الأول" in learn.text and "اقرأ أولًا" in learn.text
    file_url = re.search(r'src="(/app/rooms/[^"]+/file)"', learn.text).group(1)
    assert s.get(file_url).status_code == 200 and s.get(file_url).headers["content-type"].startswith("image/")

    # My Library shows the projects and the new assets, with links to use them elsewhere
    library = t.get("/app/library")
    assert f"/app/presentations/{deck_id}" in library.text and f"/app/learning/{cid}" in library.text
    assert f"/app/media?asset={version_id}" in library.text and "صورة مدوّرة" in library.text


def test_a_finished_lesson_offers_the_other_studios():
    from rafeeq.web.common import templates

    html = templates.get_template("partials/session.html").render(
        t=lambda k, **kw: k, u=lambda p, **q: "/app" + p + ("?" + "&".join(f"{a}={b}" for a, b in q.items()) if q else ""),
        icon=templates.env.globals["icon"], csrf="x", room={"room_id": "r_1", "name": "Physics"},
        session={"session_id": "s_9", "topic": "الضوء", "status": "done", "draft": {"slides": [{"title": "x", "bullets": []}]},
                 "selected_agents": [], "files": [], "revisions_left": 0, "memory": []})
    for path in ("/app/presentations/open?session=s_9", "/app/video?session=s_9", "/app/media?session=s_9", "/app/learning?session=s_9"):
        assert path in html, path
    assert "room=r_1" in html
    waiting = templates.get_template("partials/session.html").render(
        t=lambda k, **kw: k, u=lambda p, **q: "/app" + p, icon=templates.env.globals["icon"], csrf="x",
        room={"room_id": "r_1", "name": "Physics"},
        session={"session_id": "s_9", "topic": "x", "status": "awaiting_approval", "draft": {}, "selected_agents": [],
                 "files": [], "revisions_left": 1, "memory": []})
    assert "/app/presentations/open" not in waiting                                                     # only after approval
